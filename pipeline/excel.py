"""Excel workbook for the accountant: Resumen, Asiento, Movimientos, Excepciones, Metodología.

Sheet names and column headers are in Spanish. Amounts are written as numbers with two
decimals, converted from integer cents at the last step.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline.config import account_name

HEADER_FILL = PatternFill("solid", fgColor="E8ECEF")
BOLD = Font(bold=True)
MONEY = "#,##0.00"
RATE = "#,##0.00"

RESOLVED_ES = {"rules": "Reglas", "ai": "IA (aceptada)", "review": "Revisión humana"}
TYPE_ES = {
    "card_purchase": "Compra con tarjeta",
    "card_fee": "Comisión de tarjeta",
    "card_refund": "Reintegro de tarjeta",
    "usdc_payout": "Pago en USDC a contratista",
    "ars_topup": "Ingreso de ARS",
    "card_settlement": "Pago del resumen de tarjeta",
}


def _num(s: str | None) -> Decimal | None:
    return Decimal(s) if s not in (None, "") else None


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], money_cols=(), rate_cols=(), widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for cell in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell[0].number_format = MONEY
    for col in rate_cols:
        for cell in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell[0].number_format = RATE
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def write(close: dict, path: Path) -> None:
    s = close["summary"]
    wb = Workbook()
    wb.remove(wb.active)

    resumen = [
        ["Empresa", s["company"]],
        ["Período", s["period"]],
        ["Datos", "Empresa ficticia y movimientos sintéticos"],
        ["Movimientos", s["transactions"]],
        ["Gasto con tarjeta neto (USD)", _num(s["totals"]["card_net_usd"])],
        ["Pagos a contratistas (USDC)", _num(s["totals"]["payouts_usdc"])],
        ["Gasto total en USD", _num(s["totals"]["spend_usd"])],
        ["Gasto total en ARS", _num(s["totals"]["spend_ars"])],
        ["Resueltos por reglas", s["split"]["rules"]],
        ["Resueltos por IA (aceptados sobre el umbral)", s["split"]["ai"]],
        ["Enviados a revisión humana", s["split"]["review"]],
        ["Umbral de confianza de la IA", s["threshold"] or "IA no ejecutada"],
        ["Excepciones abiertas", s["exceptions"]],
        ["Asiento balanceado (Debe = Haber)", "Sí" if close["entry_totals"]["balanced"] else "No"],
        ["Tipo de cambio de cierre (oficial vendedor)", _num(close["balances"]["closing_rate"])],
        ["Tipo de cambio de cierre USDC (dólar cripto vendedor)", _num(close["balances"]["usdc_closing_rate"])],
    ]
    ws = _sheet(wb, "Resumen", ["Concepto", "Valor"], resumen, widths=[48, 28])
    for row in (6, 7, 8, 9):
        ws.cell(row=row, column=2).number_format = MONEY

    entry_rows = [[r["account"], r["name"], _num(r["debe"]), _num(r["haber"])] for r in close["entry"]]
    entry_rows.append(["", "Totales", _num(close["entry_totals"]["debe"]), _num(close["entry_totals"]["haber"])])
    ws = _sheet(
        wb,
        "Asiento",
        ["Cuenta", "Nombre", "Debe (ARS)", "Haber (ARS)"],
        entry_rows,
        money_cols=(3, 4),
        widths=[10, 44, 20, 20],
    )
    for c in ws[ws.max_row]:
        c.font = BOLD

    mov = []
    for t in close["transactions"]:
        mov.append(
            [
                t["id"],
                t["posted_at"],
                TYPE_ES[t["type"]],
                t["descriptor"],
                t["currency"],
                _num(t["amount"]),
                _num(t["rate"]),
                _num(t["ars"]),
                t["posted_account"],
                account_name(t["posted_account"]) if t["posted_account"] else "",
                t["account"] or "",
                RESOLVED_ES[t["resolved_by"]],
                t["rule_id"] or (t["ai"]["id"] if t["ai"] else ""),
                t["ai"]["confidence"] if t["ai"] else "",
                t["status"],
                t["receipt_id"] or "",
                ", ".join(t["exceptions"]),
            ]
        )
    _sheet(
        wb,
        "Movimientos",
        [
            "ID",
            "Fecha",
            "Tipo",
            "Descriptor",
            "Moneda",
            "Importe",
            "TC",
            "Importe ARS",
            "Cuenta imputada",
            "Nombre cuenta",
            "Cuenta propuesta",
            "Resuelto por",
            "Regla o sugerencia",
            "Confianza IA",
            "Estado",
            "Comprobante",
            "Excepciones",
        ],
        mov,
        money_cols=(6, 8),
        rate_cols=(7,),
        widths=[9, 11, 22, 30, 7, 13, 10, 16, 10, 34, 10, 16, 14, 10, 15, 11, 30],
    )

    exc = []
    for e in close["exceptions"]:
        exc.append(
            [
                e["id"],
                e["type"],
                e.get("txn_id") or "",
                e.get("receipt_id") or "",
                e.get("facts", ""),
                e.get("explanation", ""),
                e.get("explanation_source", ""),
            ]
        )
    _sheet(
        wb,
        "Excepciones",
        ["ID", "Tipo", "Movimiento", "Comprobante", "Hechos", "Explicación", "Fuente"],
        exc,
        widths=[8, 24, 11, 11, 60, 70, 14],
    )

    met = [[m["topic"], m["label"], m["detail"]] for m in close["methodology"]]
    ws = _sheet(wb, "Metodología", ["Tema", "Etiqueta", "Detalle"], met, widths=[32, 14, 110])
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
